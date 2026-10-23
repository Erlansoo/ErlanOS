#!/usr/bin/env python3
"""D0 host export bridge. Reads data emitted by ErlanOS; never executes it.

The native service owns content, revisions and layout. This temporary bridge
packages its snapshot for existing desktop applications until filesystem and
native PDF/ZIP writers exist in ErlanOS.
"""
from __future__ import annotations
import argparse
from dataclasses import asdict, dataclass
import hashlib
import json
from pathlib import Path
import re
import struct
import xml.etree.ElementTree as ET
import zipfile

ROOT = Path(__file__).resolve().parents[1]
FONTS = ROOT / 'assets/fonts'
FONT_NAMES = ('DejaVuSerif', 'DejaVuSans')
FONT_FAMILIES = ('DejaVu Serif', 'DejaVu Sans')
UM_TO_PT = 72 / 25400

@dataclass(frozen=True)
class Document:
    revision: int
    width: int
    height: int
    top: int
    bottom: int
    left: int
    right: int
    mirror: int
    font: int
    size10: int
    title: str
    text: str

def checksum(data: bytes) -> int:
    value = 14695981039346656037
    for byte in data:
        value = ((value ^ byte) * 1099511628211) & ((1 << 64) - 1)
    return value

def decode(data: bytes) -> Document:
    if not 224 <= len(data) <= 16608:
        raise ValueError('ELD1 snapshot length out of range')
    magic, rev, w, h, top, bottom, left, right, mirror, font, size, n, caret, page, title_n, editing = struct.unpack_from('<16Q', data)
    if magic != 0x31444c45 or not rev or n != len(data)-224 or not 1 <= title_n <= 96:
        raise ValueError('Invalid ELD1 header/length')
    if not (100000 <= w <= 420000 and 100000 <= h <= 420000 and
            all(m <= 150000 for m in (top,bottom,left,right)) and
            left+right+30000 <= w and top+bottom+40000 <= h and
            mirror <= 1 and font <= 1 and 80 <= size <= 240 and size%5==0 and
            caret <= n and 1 <= page <= 32 and editing <= 1):
        raise ValueError('Invalid ELD1 geometry or state')
    title = data[128:128+title_n].decode('utf-8', errors='strict')
    text = data[224:].decode('utf-8', errors='strict')
    data[224:224+caret].decode('utf-8', errors='strict')
    for value, controls in ((title,''),(text,'\n\f')):
        if any((ord(c)<32 and c not in controls) or 127<=ord(c)<160 for c in value):
            raise ValueError('Control character in document')
    return Document(rev,w,h,top,bottom,left,right,mirror,font,size,title,text)

def frames(serial: str) -> list[bytes]:
    """Only complete, integrity-checked frames; a partial trailing frame is pending.

    A completed malformed frame raises instead of silently exporting an older one.
    FNV detects corruption, not authenticity; the serial log must be trusted.
    """
    result, chunks, expected = [], None, None
    for raw in serial.splitlines(keepends=True):
        if not raw.endswith(('\n','\r')):
            break  # the live serial writer may be halfway through a hex line
        line = raw.removeprefix('[system] ').strip()
        if line == 'doc: export begin ELD1':
            chunks, expected = [], None
        elif chunks is not None:
            if line == 'doc: export end ELD1':
                data = bytes.fromhex(''.join(chunks))
                if expected is None or checksum(data) != expected:
                    raise ValueError('ELD1 export checksum mismatch')
                decode(data)
                result.append(data)
                chunks = None
            elif line.startswith('doc: export hash='):
                if expected is not None or not line[17:].isdecimal():
                    raise ValueError('Invalid ELD1 checksum')
                expected = int(line[17:])
            elif re.fullmatch(r'[0-9a-f]{2,128}', line) and len(line)%2==0 and expected is None:
                chunks.append(line)
                if sum(map(len,chunks)) > 33216:
                    raise ValueError('ELD1 export too large')
            elif line.startswith('[model]') or line.startswith('[halo]') or re.fullmatch(r'\[pit\] \d+ s',line):
                continue
            else:
                raise ValueError('Invalid line inside ELD1 export')
    return result

def font_metrics(font: int) -> list[int]:
    name = FONT_NAMES[font]
    entry = json.loads((FONTS/'metrics.json').read_text())[name]
    if hashlib.sha256((FONTS/(name+'.ttf')).read_bytes()).hexdigest() != entry['sha256']:
        raise ValueError('Font changed; regenerate native atlas and metrics first')
    return entry['advances_2048']

def layout(doc: Document) -> tuple[list[dict], int]:
    """Use the native integer advances and wrapping for the fixed-layout PDF.

    DOCX/ODT carry semantic paragraphs and may reflow in the receiving application.
    """
    metrics = font_metrics(doc.font)
    def advance(c, size):
        code = ord(c) if 32 <= ord(c) <= 255 else ord('?')
        return metrics[code-32]*size*25400//1474560
    def left(page):
        return doc.right if doc.mirror and page%2==0 else doc.left
    width, bottom, page, y = doc.width-doc.left-doc.right, doc.height-doc.bottom, 1, doc.top
    lines = []
    for text, size in ((doc.title,180),(doc.text,doc.size10)):
        height, ascent = size*25400*135//72000, size*25400//720
        start = 0
        while True:
            if y+height > bottom:
                page, y = page+1, doc.top
            if page>32 or len(lines)>=1024:
                raise ValueError('Native page/line limit exceeded')
            if start<len(text) and text[start]=='\f':
                page, y, start = page+1, doc.top, start+1
                continue
            p, used, space, nbytes = start, 0, 0, 0
            while p<len(text) and text[p] not in '\n\f':
                step, size_bytes = advance(text[p],size), len(text[p].encode())
                if (used+step>width or nbytes+size_bytes>224) and p>start:
                    break
                used, nbytes, p = used+step, nbytes+size_bytes, p+1
                if text[p-1]==' ': space=p
            end = space if p<len(text) and text[p] not in '\n\f' and space>start else p
            value = text[start:end]
            lines.append(dict(page=page,x=left(page),y=y+ascent,size10=size,text=value,
                              advances=[advance(c,size) for c in value]))
            y, start = y+height, end
            if start>=len(text): break
            if text[start]=='\n': start+=1
            else:
                while start<len(text) and text[start]==' ': start+=1
        if size==180 and text==doc.title: y+=5000
    return lines,page

def export_pdf(doc: Document, path: Path):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    name=FONT_NAMES[doc.font]
    pdfmetrics.registerFont(TTFont(name,str(FONTS/(name+'.ttf'))))
    lines,pages=layout(doc)
    c=canvas.Canvas(str(path),pagesize=(doc.width*UM_TO_PT,doc.height*UM_TO_PT),pageCompression=1)
    c.setTitle(doc.title); c.setAuthor('ErlanOS Documentos')
    for page in range(1,pages+1):
        for line in (l for l in lines if l['page']==page):
            x=line['x']; y=(doc.height-line['y'])*UM_TO_PT
            c.setFont(name,line['size10']/10)
            # Explicit advances preserve the native page geometry without kerning drift.
            for ch,step in zip(line['text'],line['advances']):
                c.drawString(x*UM_TO_PT,y,ch)
                x+=step
        c.showPage()
    c.save()

def export_docx(doc: Document, path: Path):
    from docx import Document as WordDocument
    from docx.shared import Pt, Emu, RGBColor
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    out=WordDocument(); sec=out.sections[0]
    sec.page_width,sec.page_height=Emu(doc.width*36),Emu(doc.height*36)
    for key in ('top','bottom','left','right'):
        setattr(sec,key+'_margin',Emu(getattr(doc,key)*36))
    sec.header_distance=sec.footer_distance=Emu(0)
    if doc.mirror: out.settings.element.append(OxmlElement('w:mirrorMargins'))
    for sty,size in (('Normal',doc.size10/10),('Title',18)):
        s=out.styles[sty]; s.font.name=FONT_FAMILIES[doc.font]; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
        s.font.bold=False; s.font.italic=False
        s.paragraph_format.space_before=Pt(0); s.paragraph_format.space_after=Pt(0)
        s.paragraph_format.line_spacing=Pt(size*1.35)
        s.paragraph_format.widow_control=False
        s.paragraph_format.keep_with_next=False
        fonts=s.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ('ascii','hAnsi','eastAsia','cs'): fonts.set(qn('w:'+key),FONT_FAMILIES[doc.font])
        for key in list(fonts.attrib):
            if key.endswith('Theme'): del fonts.attrib[key]
    title=out.add_paragraph(doc.title,'Title'); title.paragraph_format.space_after=Emu(5000*36)
    # Bootstrap templates may carry a themed blue rule in the Title style.
    # D0 has no paragraph borders; do not add styling absent from the native page.
    for border in list(out.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for text in doc.text.split('\n'):
        p=out.add_paragraph(style='Normal')
        for i,part in enumerate(text.split('\f')):
            if i: p.add_run().add_break(WD_BREAK.PAGE)
            p.add_run(part)
    out.core_properties.title=doc.title; out.core_properties.author='ErlanOS Documentos'
    out.save(path)

NS={p:'urn:oasis:names:tc:opendocument:xmlns:'+tail+':1.0' for p,tail in
    (('office','office'),('style','style'),('text','text'),('fo','xsl-fo-compatible'),
     ('svg','svg-compatible'),('manifest','manifest'),('meta','meta'))}
NS['dc']='http://purl.org/dc/elements/1.1/'
for prefix,uri in NS.items(): ET.register_namespace(prefix,uri)
def tag(name):
    prefix,local=name.split(':'); return '{'+NS[prefix]+'}'+local
def element(parent,name,attrs=None):
    attrs={tag(k):str(v) for k,v in (attrs or {}).items()}
    return ET.SubElement(parent,tag(name),attrs) if parent is not None else ET.Element(tag(name),attrs)
def xml(root): return ET.tostring(root,encoding='utf-8',xml_declaration=True)

def export_odt(doc: Document, path: Path):
    styles=element(None,'office:document-styles',{'office:version':'1.3'})
    faces=element(styles,'office:font-face-decls')
    element(faces,'style:font-face',{'style:name':'DocumentFont','svg:font-family':FONT_FAMILIES[doc.font]})
    named=element(styles,'office:styles')
    for name,size in (('Body',doc.size10/10),('Title',18),('PageBreak',doc.size10/10)):
        attrs={'style:name':name,'style:family':'paragraph'}
        if name=='Title': attrs['style:master-page-name']='Standard'
        if name=='PageBreak': attrs['style:parent-style-name']='Body'
        s=element(named,'style:style',attrs)
        pp={'fo:margin-top':'0cm','fo:margin-bottom':'0.5cm' if name=='Title' else '0cm',
            'fo:line-height':f'{size*1.35:g}pt','fo:widows':'1','fo:orphans':'1'}
        if name=='PageBreak': pp['fo:break-before']='page'
        element(s,'style:paragraph-properties',pp)
        element(s,'style:text-properties',{'style:font-name':'DocumentFont','fo:font-size':f'{size:g}pt',
                                          'fo:color':'#000000','fo:font-weight':'normal'})
    auto=element(styles,'office:automatic-styles')
    pg=element(auto,'style:page-layout',{'style:name':'Page','style:page-usage':'mirrored' if doc.mirror else 'all'})
    props={'fo:page-width':f'{doc.width/10000:g}cm','fo:page-height':f'{doc.height/10000:g}cm',
           'style:print-orientation':'portrait' if doc.height>=doc.width else 'landscape'}
    for key in ('top','bottom','left','right'): props['fo:margin-'+key]=f'{getattr(doc,key)/10000:g}cm'
    element(pg,'style:page-layout-properties',props)
    masters=element(styles,'office:master-styles')
    element(masters,'style:master-page',{'style:name':'Standard','style:page-layout-name':'Page'})
    content=element(None,'office:document-content',{'office:version':'1.3'})
    body=element(element(content,'office:body'),'office:text')
    element(body,'text:p',{'text:style-name':'Title'}).text=doc.title
    for para in doc.text.split('\n'):
        for i,part in enumerate(para.split('\f')):
            p=element(body,'text:p',{'text:style-name':'PageBreak' if i else 'Body'})
            # Explicit space elements preserve consecutive and leading spaces.
            previous=None
            for token in re.split('( +)',part):
                if token.startswith(' '): previous=element(p,'text:s',{'text:c':len(token)})
                elif previous is None: p.text=(p.text or '')+token
                else: previous.tail=(previous.tail or '')+token
    meta=element(None,'office:document-meta',{'office:version':'1.3'})
    element(element(meta,'office:meta'),'dc:title').text=doc.title
    manifest=element(None,'manifest:manifest',{'manifest:version':'1.3'})
    mime='application/vnd.oasis.opendocument.text'
    for name,kind in (('/',mime),('content.xml','text/xml'),('styles.xml','text/xml'),('meta.xml','text/xml')):
        element(manifest,'manifest:file-entry',{'manifest:full-path':name,'manifest:media-type':kind})
    with zipfile.ZipFile(path,'w',compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr('mimetype',mime,compress_type=zipfile.ZIP_STORED)
        for name,value in (('content.xml',content),('styles.xml',styles),('meta.xml',meta),('META-INF/manifest.xml',manifest)):
            z.writestr(name,xml(value))

def export_all(data: bytes, directory: Path, name='documento') -> dict:
    if not re.fullmatch('[A-Za-z0-9_-]{1,80}',name): raise ValueError('Unsafe output basename')
    doc=decode(data); lines,pages=layout(doc)
    directory.mkdir(parents=True,exist_ok=True)
    (directory/(name+'.eld')).write_bytes(data)
    description=asdict(doc)|{'format':'ELD1','units':'micrometres','native_pages':pages,
        'font_family':FONT_FAMILIES[doc.font],'fnv1a64':checksum(data),
        'export_bridge':'host D0; PDF uses native line positions; DOCX/ODT may reflow'}
    (directory/(name+'.json')).write_text(json.dumps(description,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    export_pdf(doc,directory/(name+'.pdf')); export_docx(doc,directory/(name+'.docx')); export_odt(doc,directory/(name+'.odt'))
    return description

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input',type=Path,help='serial.log or .eld snapshot')
    parser.add_argument('--out',type=Path,required=True)
    parser.add_argument('--name',default='documento')
    args=parser.parse_args()
    data=args.input.read_bytes()
    if args.input.suffix.lower()!='.eld':
        found=frames(data.decode('utf-8',errors='replace'))
        if not found: parser.error('No complete ELD1 export in serial log')
        data=found[-1]
    result=export_all(data,args.out,args.name)
    print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=='__main__': main()
